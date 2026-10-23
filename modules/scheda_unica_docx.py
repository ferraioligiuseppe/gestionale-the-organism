# -*- coding: utf-8 -*-
"""
SCHEDA CLINICA COMPLETA PNEV — un unico Word da stampare e compilare a mano.

Riunisce in un solo documento tutte le parti della visita: anagrafica,
anamnesi, Profilo Neuro-Evolutivo (tutte le prove del protocollo, prese da
inpp/protocollo.py: se cambia il protocollo, cambia anche la scheda),
valutazione visiva, uditiva, apprendimenti e conclusioni.

Uso nel gestionale:
    from modules.scheda_unica_docx import genera_scheda_unica, MIME_DOCX
    st.download_button("Scheda completa", genera_scheda_unica(), "Scheda_PNEV.docx", MIME_DOCX)
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (disponibile se serve)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

VERDE = RGBColor(0x1D, 0x6B, 0x44)
GRIGIO = RGBColor(0x6B, 0x6B, 0x66)
VERDE_CHIARO = "E8F2EC"
BOX = "☐"

SEDI = ("Pagani (SA) · Via De Rosa 46   |   Sant'Agnello (NA) · Via Balsamo 19   |   "
        "Avellino · SociaLab, Via Renzi 40")
CONTATTI = "WhatsApp 391 3598767   ·   apstheorganism@gmail.com   ·   www.pnev.it"


# ════════════════════════════════════════════════════════════════════
#  MATTONI DI IMPAGINAZIONE
# ════════════════════════════════════════════════════════════════════

def _shade(cell, colore: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), colore)
    tcPr.append(shd)


def _bordo_sotto(par, colore="1D6B44", spessore=8) -> None:
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(spessore))
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), colore)
    bdr.append(b)
    pPr.append(bdr)


def _no_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def _ripeti_intestazione(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    trPr.append(el)


def _numero_pagina(par) -> None:
    for tipo, testo in (("begin", None), (None, "PAGE"), ("end", None)):
        run = par.add_run()
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = testo
            run._r.append(it)


def _testo(par, testo, *, bold=False, size=None, color=None, italic=False):
    r = par.add_run(testo)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


class Scheda:
    def __init__(self):
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(10)
        st.paragraph_format.space_after = Pt(2)
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(1.6)
        sec.top_margin, sec.bottom_margin = Cm(1.4), Cm(1.4)
        self.larghezza = sec.page_width - sec.left_margin - sec.right_margin
        self._intestazione(sec)

    # ── intestazione e piè di pagina ──
    def _intestazione(self, sec):
        h = sec.header.paragraphs[0]
        _testo(h, "STUDIO THE ORGANISM", bold=True, size=11, color=VERDE)
        _testo(h, "   ·   Metodo PNEV — Psico-Neuro-Evolutivo", size=9, color=GRIGIO)
        _bordo_sotto(h)
        f = sec.footer.paragraphs[0]
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _testo(f, SEDI, size=7, color=GRIGIO)
        f2 = sec.footer.add_paragraph()
        f2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _testo(f2, CONTATTI + "   ·   pag. ", size=7, color=GRIGIO)
        _numero_pagina(f2)
        for r in f2.runs[-3:]:
            r.font.size = Pt(7)

    # ── elementi ──
    def titolo_parte(self, numero, testo, nota=None, nuova_pagina=True):
        if nuova_pagina:
            self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        _testo(p, f"{numero}  ", bold=True, size=15, color=VERDE)
        _testo(p, testo.upper(), bold=True, size=13, color=VERDE)
        _bordo_sotto(p, spessore=6)
        if nota:
            n = self.doc.add_paragraph()
            _testo(n, nota, italic=True, size=8.5, color=GRIGIO)

    def sottotitolo(self, testo):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        _testo(p, testo, bold=True, size=10.5, color=VERDE)

    def campi(self, coppie, colonne=2):
        """Griglia di campi «Etichetta: ________». Una coppia può essere una
        stringa (campo con riga) o (etichetta, opzioni) per le caselle ☐."""
        righe = (len(coppie) + colonne - 1) // colonne
        t = self.doc.add_table(rows=righe, cols=colonne)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        w = int(self.larghezza / colonne)
        for i, campo in enumerate(coppie):
            cell = t.cell(i // colonne, i % colonne)
            cell.width = w
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(3)
            if isinstance(campo, tuple):
                etichetta, opzioni = campo
                _testo(p, f"{etichetta}: ", bold=True, size=9)
                _testo(p, "   ".join(f"{BOX} {o}" for o in opzioni), size=9)
            else:
                _testo(p, f"{campo}: ", bold=True, size=9)
                _testo(p, "_" * max(8, 46 // colonne * 2 - len(campo)), size=9, color=GRIGIO)
        for row in t.rows:
            _no_split(row)

    def righe(self, n, etichetta=None):
        if etichetta:
            p = self.doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            _testo(p, etichetta, bold=True, size=9)
        for _ in range(n):
            # riga da scrivere: tabulazione sottolineata fino al margine destro
            # (i bordi di paragrafi uguali e vicini Word li fonde in uno solo)
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.tab_stops.add_tab_stop(self.larghezza)
            r = p.add_run("\t")
            r.underline = True
            r.font.color.rgb = RGBColor(0xBB, 0xBB, 0xB5)

    def tabella(self, intestazioni, righe, larghezze_cm, *, prima_colonna_bold=True,
                righe_totale=()):
        t = self.doc.add_table(rows=1, cols=len(intestazioni))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        # larghezze anche sulla griglia della tabella (Word e LibreOffice)
        grid = t._tbl.tblGrid
        for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(Cm(larghezze_cm[j]).twips)))
        for j, col in enumerate(t.columns):
            col.width = Cm(larghezze_cm[j])
        for j, h in enumerate(intestazioni):
            c = t.rows[0].cells[j]
            c.width = Cm(larghezze_cm[j])
            _shade(c, VERDE_CHIARO)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
            _testo(p, h, bold=True, size=8.5, color=VERDE)
        _ripeti_intestazione(t.rows[0])
        for i, valori in enumerate(righe):
            row = t.add_row()
            _no_split(row)
            tot = i in righe_totale
            # riga-titolo di gruppo: testo solo nella prima cella → cella unica
            if tot and all(v == "" for v in valori[1:]):
                c = row.cells[0].merge(row.cells[-1])
                _shade(c, "F4F4F1")
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.keep_with_next = True
                _testo(p, valori[0], bold=True, size=8, color=VERDE)
                continue
            for j, v in enumerate(valori):
                c = row.cells[j]
                c.width = Cm(larghezze_cm[j])
                if tot:
                    _shade(c, "F4F4F1")
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                if j and v in (BOX,) or (j and isinstance(v, str) and v.startswith(BOX) and len(v) < 4):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _testo(p, v, bold=(j == 0 and prima_colonna_bold) or tot, size=8.5)
        return t

    def salva(self) -> bytes:
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


# ════════════════════════════════════════════════════════════════════
#  CONTENUTO
# ════════════════════════════════════════════════════════════════════

def _copertina(s: Scheda):
    p = s.doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    _testo(p, "Scheda clinica completa", bold=True, size=22, color=VERDE)
    p = s.doc.add_paragraph()
    _testo(p, "Valutazione PNEV — anamnesi, profilo neuro-evolutivo, visione, udito, apprendimenti",
           size=10.5, color=GRIGIO)
    s.doc.add_paragraph()
    s.campi(["Paziente", "Data della visita", "Professionista", "Sede",
             ("Fascia d'età", ["0–4", "5–16", "17–64", "over 65"]),
             ("Prima visita", ["sì", "controllo"])], colonne=2)
    s.sottotitolo("Percorso di valutazione (spunta le parti svolte)")
    s.tabella(
        ["Parte", "Svolta", "Data", "Note"],
        [["1  Anagrafica", BOX, "", ""], ["2  Anamnesi PNEV", BOX, "", ""],
         ["3  Profilo Neuro-Evolutivo", BOX, "", ""], ["4  Valutazione visiva", BOX, "", ""],
         ["5  Valutazione uditiva", BOX, "", ""], ["6  Apprendimenti (dai 6 anni)", BOX, "", ""],
         ["7  Sintesi e percorso", BOX, "", ""]],
        [6.2, 1.6, 2.8, 7.2])


def _anagrafica(s: Scheda):
    s.titolo_parte("1", "Anagrafica")
    s.sottotitolo("Paziente")
    s.campi(["Cognome", "Nome", "Data di nascita", ("Sesso", ["M", "F"]),
             "Comune di nascita (prov.)", "Codice fiscale", "Indirizzo (via, civico)",
             "CAP e città", "Cellulare", "Telefono fisso", "Email",
             "Medico / pediatra"])
    s.sottotitolo("Scuola (se in età scolare)")
    s.campi(["Scuola", "Classe", "Insegnante di riferimento", ("Sostegno / PDP / PEI", ["no", "PDP", "PEI"])])
    s.sottotitolo("Genitori / tutore (se minore)")
    s.campi(["Madre — nome", "Madre — cellulare", "Padre — nome", "Padre — cellulare",
             "Tutore (se diverso)", "Email per i documenti"])
    s.sottotitolo("Invio e motivo")
    s.campi(["Inviato da", ("Canale", ["medico", "scuola", "passaparola", "web", "altro"])])
    s.righe(3, "Motivo della visita (con le parole del paziente / dei genitori)")
    s.sottotitolo("Consensi")
    s.campi([("Privacy (GDPR)", ["firmato"]), ("Consenso informato", ["firmato"]),
             ("Video a scopo clinico", ["sì", "no"]), ("Comunicazioni via WhatsApp", ["sì", "no"])])


def _anamnesi(s: Scheda):
    s.titolo_parte("2", "Anamnesi PNEV",
                   "Per gli adulti compilare le parti pertinenti; le informazioni sulla prima "
                   "infanzia si raccolgono quando note.")
    s.sottotitolo("Gravidanza")
    s.campi([("Termine", ["a termine", "pretermine", "post-termine"]), "Settimane",
             ("Tipo", ["singola", "gemellare", "PMA"]), ("Movimenti fetali", ["regolari", "scarsi", "eccessivi"]),
             ("Controlli prenatali", ["regolari", "irregolari"]),
             ("Andamento", ["fisiologico", "con complicanze"])])
    s.righe(1, "Complicanze, farmaci, stress in gravidanza")
    s.sottotitolo("Parto")
    s.campi([("Tipo", ["spontaneo", "indotto", "cesareo", "ventosa/forcipe"]),
             ("Presentazione", ["cefalica", "podalica", "altro"]), ("Pianto immediato", ["sì", "no"]),
             ("Contatto pelle-pelle", ["sì", "no"]), ("Suzione immediata", ["sì", "no"]),
             "Peso alla nascita / Apgar"])
    s.sottotitolo("Periodo neonatale")
    s.campi([("Tono muscolare", ["normale", "ipotono", "ipertono"]), ("Pianto", ["normale", "debole", "eccessivo"]),
             ("Moro", ["presente", "assente", "asimmetrico"]), ("Prensione", ["presente", "debole"]),
             ("Suzione", ["valida", "debole"]), ("Rooting", ["presente", "assente"]),
             ("Ittero / incubatrice", ["no", "sì"]), ("Coliche / reflusso", ["no", "sì"])])
    s.sottotitolo("Alimentazione")
    s.campi([("Allattamento", ["seno", "artificiale", "misto"]), "Durata",
             ("Qualità", ["valida", "difficoltosa"]), ("Svezzamento / masticazione", ["regolare", "difficoltoso"]),
             ("Selettività alimentare", ["no", "sì"]), ("Respirazione", ["nasale", "orale", "mista"])])
    s.sottotitolo("Sviluppo motorio")
    s.campi([("Simmetria posturale", ["sì", "no"]), ("Tilt test (raddrizzamento)", ["presente", "debole"]),
             ("ATNR (schermitore)", ["no", "sì"]), ("Gattonamento", ["sì", "no", "breve"]),
             ("Schema gattonamento", ["crociato", "omolaterale", "strisciato"]),
             ("Simmetria gattonamento", ["sì", "no"]), "Età primi passi",
             ("Appoggio plantare", ["normale", "punte", "piatto"]), ("Base d'appoggio", ["normale", "allargata"]),
             ("Cadute frequenti", ["no", "sì"]), ("Riflessi protettivi", ["presenti", "deboli"]),
             ("Bicicletta / nuoto", ["sì", "no"])])
    s.sottotitolo("Linguaggio e sviluppo sensoriale")
    s.campi(["Prime parole (età)", "Prime frasi (età)", ("Risposta al nome", ["sì", "no", "incostante"]),
             ("Contatto oculare", ["buono", "sfuggente"]), ("Risposta ai suoni", ["adeguata", "ipo", "iper"]),
             ("Otiti ricorrenti", ["no", "sì"]), ("Sensibilità tattile", ["adeguata", "ipo", "iper"]),
             ("Sonno", ["regolare", "difficoltoso", "russa"])])
    s.sottotitolo("Segnali di attenzione")
    s.tabella(["Area", "No", "Sì", "Note"],
              [["Attenzione / iperattività", BOX, BOX, ""], ["Emotività / ansia", BOX, BOX, ""],
               ["Coordinazione / goffaggine", BOX, BOX, ""], ["Lettura / scrittura", BOX, BOX, ""],
               ["Mal d'auto / di mare", BOX, BOX, ""], ["Mal di testa / affaticamento visivo", BOX, BOX, ""],
               ["Enuresi oltre i 5 anni", BOX, BOX, ""], ["Allergie / asma / eczema", BOX, BOX, ""]],
              [6.4, 1.3, 1.3, 8.8])
    s.sottotitolo("Famiglia e salute")
    s.campi([("Fratelli/sorelle con difficoltà", ["no", "sì"]), "Lingua parlata in casa",
             "Familiarità (DSA, ADHD, visione…)", "Terapie in corso / precedenti",
             "Patologie / interventi", "Farmaci"])
    s.righe(3, "Osservazioni")


def _valore_cella(prova):
    sc = prova.get("scoring", "0-4")
    if sc == "lateralita":
        return f"{BOX} Dx   {BOX} Sx   {BOX} Misto"
    if sc == "si_no":
        return f"{BOX} Sì   {BOX} No"
    if sc == "scelta":
        return "   ".join(f"{BOX} {o}" for o in prova.get("opzioni", []))
    if sc == "numerico":
        return "n. ________"
    return ""


def _profilo_neuro(s: Scheda):
    s.titolo_parte("3", "Profilo Neuro-Evolutivo PNEV",
                   "Valutazione della maturità neuro-psico-motoria — segni neurologici minori e "
                   "integrazione sensomotoria. Punteggio 0 = nessuna difficoltà · 1 = 25% · 2 = 50% · "
                   "3 = 75% · 4 = completa (100%). Prove in parte adattate dal protocollo INPP (rev. 01/22).")
    try:
        from .inpp.protocollo import PROTOCOLLO_INPP
    except ImportError:
        from modules.inpp.protocollo import PROTOCOLLO_INPP
    larg = [6.6, 0.75, 0.75, 0.75, 0.75, 0.75, 7.45]
    for n, sezione in enumerate(PROTOCOLLO_INPP, 1):
        s.sottotitolo(f"3.{n}  {sezione['label']}")
        righe, n_punteggio, totali = [], 0, []
        for gruppo in sezione["gruppi"]:
            if gruppo.get("label") and len(sezione["gruppi"]) > 1:
                righe.append([gruppo["label"].upper(), "", "", "", "", "", ""])
                totali.append(len(righe) - 1)
            for prova in gruppo["prove"]:
                sc = prova.get("scoring", "0-4")
                if sc == "0-4":
                    n_punteggio += 1
                    righe.append([prova["label"], BOX, BOX, BOX, BOX, BOX, ""])
                elif sc == "testo":
                    righe.append([prova["label"], "", "", "", "", "", ""])
                else:
                    righe.append([prova["label"], "", "", "", "", "", _valore_cella(prova)])
        if n_punteggio and not sezione.get("no_total"):
            righe.append([f"Totale sezione  ( ____ / {n_punteggio * 4} )", "", "", "", "", "",
                          "Indice di immaturità: ______ %"])
            totali.append(len(righe) - 1)
        s.tabella(["Prova", "0", "1", "2", "3", "4", "Note / risposta"], righe, larg,
                  prima_colonna_bold=False, righe_totale=totali)
    s.sottotitolo("Sintesi del profilo")
    s.campi([("Livello di immaturità", ["assente", "lieve", "moderato", "marcato"]),
             "Riflessi da integrare per primi"])
    s.righe(3, "Osservazioni cliniche")


def _visiva(s: Scheda):
    s.titolo_parte("4", "Valutazione visiva")
    s.sottotitolo("Acuità visiva")
    s.tabella(["", "OD", "OS", "OO", "Note"],
              [["Lontano naturale", "", "", "", ""], ["Lontano con correzione", "", "", "", ""],
               ["Vicino naturale", "", "", "", ""], ["Vicino con correzione", "", "", "", ""]],
              [5.2, 2.4, 2.4, 2.4, 5.4])
    s.sottotitolo("Refrazione")
    s.tabella(["", "Sfera", "Cilindro", "Asse", "Add.", "Visus"],
              [["Abituale OD", "", "", "", "", ""], ["Abituale OS", "", "", "", ""] + [""],
               ["Oggettiva OD", "", "", "", "", ""], ["Oggettiva OS", "", "", "", "", ""],
               ["Soggettiva OD", "", "", "", "", ""], ["Soggettiva OS", "", "", "", "", ""]],
              [5.2, 2.5, 2.5, 2.5, 2.5, 2.6])
    s.sottotitolo("Visione binoculare e motilità")
    s.campi(["Cover test lontano", "Cover test vicino", "PPC rottura / recupero (cm)",
             "Ampiezza accomodativa OD / OS", "Flessibilità accomodativa (cpm)", "Stereopsi (″)",
             ("Motilità oculare", ["fluida", "a scatti", "perdita di fissazione"]),
             ("Saccadi", ["adeguate", "imprecise"]), ("Dominanza oculare", ["Dx", "Sx"]),
             ("Postura / distanza di lavoro", ["adeguata", "ravvicinata", "inclinata"])])
    s.sottotitolo("Test di efficienza visiva")
    s.tabella(["Test", "Risultato", "Norma / classe", "Note"],
              [["DEM — tempo verticale (A+B)", "", "", ""], ["DEM — tempo orizzontale (C)", "", "", ""],
               ["DEM — errori · rapporto · tipo", "", "", ""],
               ["Getman — punteggio /12", "", "", ""], ["Groffman — punteggio totale", "", "", ""],
               ["Visuo-percettivo (TVPS / altro)", "", "", ""]],
              [6.4, 3.4, 3.4, 4.6])
    s.righe(3, "Osservazioni visive")


def _uditiva(s: Scheda):
    s.titolo_parte("5", "Valutazione uditiva")
    s.sottotitolo("Audiometria funzionale (dB)")
    freq = ["125", "250", "500", "1000", "2000", "3000", "4000", "6000", "8000"]
    s.tabella(["Hz"] + freq,
              [["OD via aerea"] + [""] * 9, ["OS via aerea"] + [""] * 9,
               ["OD via ossea"] + [""] * 9, ["OS via ossea"] + [""] * 9],
              [3.6] + [1.58] * 9)
    s.sottotitolo("Elaborazione uditiva")
    s.campi([("Lateralità uditiva", ["Dx", "Sx", "non definita"]), "Selettività (curva)",
             ("Iper-sensibilità ai suoni", ["no", "sì"]), ("Ascolto in ambiente rumoroso", ["adeguato", "difficile"]),
             "Fisher — punteggio", "Discriminazione / ascolto dicotico"])
    s.righe(3, "Osservazioni uditive")


def _apprendimenti(s: Scheda):
    s.titolo_parte("6", "Apprendimenti", "Dai 6 anni. Indicare le prove usate e i punteggi rispetto alla norma.")
    s.tabella(["Area", "Prova usata", "Velocità / punteggio", "Errori", "Fascia"],
              [["Lettura di brano", "", "", "", f"{BOX} norma {BOX} att. {BOX} RII"],
               ["Lettura di parole / non parole", "", "", "", f"{BOX} norma {BOX} att. {BOX} RII"],
               ["Comprensione del testo", "", "", "", f"{BOX} norma {BOX} att. {BOX} RII"],
               ["Scrittura (dettato)", "", "", "", f"{BOX} norma {BOX} att. {BOX} RII"],
               ["Grafia / grafo-motricità", "", "", "", f"{BOX} norma {BOX} att. {BOX} RII"],
               ["Calcolo e numero", "", "", "", f"{BOX} norma {BOX} att. {BOX} RII"],
               ["Attenzione / memoria di lavoro", "", "", "", f"{BOX} norma {BOX} att. {BOX} RII"]],
              [4.6, 3.4, 3.0, 1.8, 5.0])
    s.doc.add_paragraph()
    p = s.doc.add_paragraph()
    _testo(p, "Legenda fascia: norma · att. = richiesta di attenzione · RII = richiesta di intervento immediato.",
           italic=True, size=8, color=GRIGIO)
    s.righe(3, "Osservazioni sugli apprendimenti")


def _sintesi(s: Scheda):
    s.titolo_parte("7", "Sintesi e percorso proposto")
    s.righe(3, "Punti di forza")
    s.righe(3, "Aree di fragilità")
    s.righe(4, "Ipotesi di lavoro")
    s.sottotitolo("Percorso proposto")
    s.tabella(["Intervento", "Sì", "Frequenza / durata", "Note"],
              [["Integrazione dei riflessi primitivi", BOX, "", ""],
               ["Stimolazione uditiva (MAPS)", BOX, "", ""],
               ["Optometria comportamentale / visual training", BOX, "", ""],
               ["Lavoro neuro-psico-motorio", BOX, "", ""],
               ["Lenti / prismi", BOX, "", ""],
               ["Lavoro a casa con la famiglia", BOX, "", ""],
               ["Invio ad altro specialista", BOX, "", ""],
               ["Controllo", BOX, "", ""]],
              [6.8, 1.2, 4.2, 5.6])
    s.doc.add_paragraph()
    s.campi(["Data", "Firma del professionista", "Firma del paziente / genitore", "Prossimo appuntamento"])


def genera_scheda_unica() -> bytes:
    s = Scheda()
    _copertina(s)
    _anagrafica(s)
    _anamnesi(s)
    _profilo_neuro(s)
    _visiva(s)
    _uditiva(s)
    _apprendimenti(s)
    _sintesi(s)
    return s.salva()


if __name__ == "__main__":
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else "Scheda_clinica_completa_PNEV.docx"
    with open(out, "wb") as f:
        f.write(genera_scheda_unica())
    print("scritto", out)
