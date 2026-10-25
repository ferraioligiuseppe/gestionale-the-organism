# -*- coding: utf-8 -*-
"""Stabilometria — esami con la pedana Wii (Balance Board) nella scheda paziente.

Gli esami si fanno con la pagina «Stabilometria PNEV» (file HTML aperto in
Chrome sul PC collegato alla pedana) e si caricano qui con il file
«Esporta per il gestionale» (oppure con la copia di sicurezza dell'archivio).

Qui restano salvati nel database, per paziente, con:
- dettaglio dell'esame (indici, profilo sensoriale, limiti di stabilità,
  parametri, tracciati, commento modificabile);
- confronto tra due esami (prima/dopo) e andamento nel tempo;
- referto in Word.

Strumento NON certificato come dispositivo medico: valori per confronti
nello stesso paziente nel tempo.

Tabella: stabilometria_esami (creata in automatico).
"""
from __future__ import annotations

import io
import json
import datetime as dt

import streamlit as st

PAGINA_PEDANA = "static_protocollo/Stabilometria_PNEV.html"

PROVE = {
    "OA": "Occhi aperti", "OC": "Occhi chiusi",
    "TDX": "Testa ruotata a destra", "TSX": "Testa ruotata a sinistra", "TES": "Testa estesa",
    "DC": "Doppio compito", "OPT": "Optocinetico", "INS": "Inseguimento lento", "SAC": "Saccadi",
    "LEN": "Lenti / prismi", "BIT": "Bite / cotoni", "LOS": "Limiti di stabilità",
    "MDX": "Monopodalico destro", "MSX": "Monopodalico sinistro",
    "GOA": "Gommapiuma occhi aperti", "GOC": "Gommapiuma occhi chiusi",
}
RIFERIMENTO = {"OC": "OA", "TDX": "OC", "TSX": "OC", "TES": "OC", "DC": "OA", "OPT": "OA",
               "INS": "OA", "SAC": "OA", "LEN": "OA", "BIT": "OC", "GOA": "OA", "GOC": "OC"}
COLORE = {"OA": "#2D5BA8", "OC": "#B5461F"}
VIOLA = "#7A4FA0"

# (etichetta, chiave, decimali, «più alto = peggio»)
PARAMETRI = [
    ("Area dell'ellisse 95% (mm²)", "area", 0, True),
    ("Lunghezza del percorso (mm)", "L", 0, True),
    ("Velocità media (mm/s)", "vel", 1, True),
    ("Velocità destra/sinistra (mm/s)", "velX", 1, True),
    ("Velocità avanti/dietro (mm/s)", "velY", 1, True),
    ("RMS destra/sinistra (mm)", "rmsX", 1, True),
    ("RMS avanti/dietro (mm)", "rmsY", 1, True),
    ("Escursione destra/sinistra (mm)", "rangeX", 1, True),
    ("Escursione avanti/dietro (mm)", "rangeY", 1, True),
    ("Lunghezza/superficie, LFS (1/mm)", "lfs", 2, False),
    ("Frequenza media dx/sx (Hz)", "fMediaX", 2, False),
    ("Frequenza media av/ind (Hz)", "fMediaY", 2, False),
    ("Entropia campionaria dx/sx", "sampEnX", 2, False),
    ("Entropia campionaria av/ind", "sampEnY", 2, False),
    ("Diffusione breve termine Ds (mm²/s)", "Ds", 1, False),
    ("Diffusione lungo termine Dl (mm²/s)", "Dl", 1, False),
    ("Tempo critico (s)", "tc", 2, False),
    ("Carico a destra (%)", "destra", 0, False),
    ("Carico in avanti (%)", "avanti", 0, False),
    ("Durata registrata (s)", "T", 1, False),
]
LOS_NOMI = {90: "Avanti", 45: "Avanti dx", 0: "Destra", 315: "Dietro dx",
            270: "Dietro", 225: "Dietro sx", 180: "Sinistra", 135: "Avanti sx"}


# ════════════════════════════════════════════════════════════════════
#  DATABASE
# ════════════════════════════════════════════════════════════════════

def _righe(conn, sql, params=()):
    try:
        cur = conn.cursor()
        cur.execute(sql, params)
        if cur.description is None:
            return []
        cols = [d[0] for d in cur.description]
        out = [r if isinstance(r, dict) else dict(zip(cols, r)) for r in cur.fetchall()]
        try:
            conn.commit()
        except Exception:
            pass
        return out
    except Exception:
        try:
            conn.rollback()
        except Exception:
            pass
        return []


def _esegui(conn, sql, params=()) -> str | None:
    try:
        cur = conn.cursor()
        cur.execute(sql, params)
        conn.commit()
        return None
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        return str(e)


def _assicura_tabella(conn) -> None:
    if _righe(conn, "SELECT 1 AS ok FROM information_schema.tables WHERE table_name='stabilometria_esami'"):
        return
    _esegui(conn, """
        CREATE TABLE IF NOT EXISTS stabilometria_esami (
            id BIGSERIAL PRIMARY KEY,
            paziente_id BIGINT NOT NULL,
            data_esame DATE,
            tipo TEXT,
            origine_id BIGINT,
            nome_su_file TEXT,
            payload JSONB,
            commento TEXT,
            creato_il TIMESTAMPTZ DEFAULT now(),
            UNIQUE (paziente_id, origine_id)
        )""")


def _json(v):
    if isinstance(v, (dict, list)):
        return v
    try:
        return json.loads(v or "{}")
    except Exception:
        return {}


def _esami(conn, paz_id) -> list[dict]:
    righe = _righe(conn, "SELECT id, data_esame, tipo, nome_su_file, payload, commento, creato_il "
                         "FROM stabilometria_esami WHERE paziente_id=%s ORDER BY data_esame, id",
                   (int(paz_id),))
    for r in righe:
        r["payload"] = _json(r.get("payload"))
    return righe


def _nome_paziente(conn, paz_id) -> str:
    r = _righe(conn, "SELECT cognome, nome FROM pazienti WHERE id=%s", (int(paz_id),))
    return f"{(r[0].get('cognome') or '').strip()} {(r[0].get('nome') or '').strip()}".strip() if r else ""


def _norm(s: str) -> str:
    return " ".join(sorted((s or "").upper().split()))


# ════════════════════════════════════════════════════════════════════
#  FORMATTAZIONE
# ════════════════════════════════════════════════════════════════════

def f1(v, d=1) -> str:
    try:
        if v is None:
            return "—"
        v = float(v)
        if v != v:
            return "—"
        return f"{v:.{d}f}".replace(".", ",")
    except Exception:
        return "—"


def _data(v) -> str:
    try:
        if isinstance(v, str):
            v = dt.date.fromisoformat(v[:10])
        return v.strftime("%d/%m/%Y")
    except Exception:
        return str(v or "")


def indici(prove: dict) -> list[tuple[str, str, str]]:
    out = []
    g = prove.get
    if g("OA") and g("OC"):
        out.append(("Quoziente di Romberg", f1(100 * g("OC")["area"] / g("OA")["area"], 0), "area OC / OA × 100"))
    if g("OC") and g("GOC"):
        out.append(("Quoziente plantare", f1(100 * g("GOC")["area"] / g("OC")["area"], 0), "area gommapiuma OC / OC × 100"))
    for pid, nome, base, k in (("DC", "Costo del doppio compito", "OA", "area"),
                               ("LEN", "Effetto lenti/prismi", "OA", "area"),
                               ("OPT", "Effetto optocinetico", "OA", "vel")):
        if g(base) and g(pid):
            pc = 100 * (g(pid)[k] / g(base)[k] - 1)
            out.append((nome, ("+" if pc >= 0 else "") + f1(pc, 0) + " %", f"{'area' if k=='area' else 'velocità'} rispetto a {base}"))
    if g("MDX") and g("MSX"):
        out.append(("Monopodalico dx / sx", f1(g("MDX")["T"]) + " / " + f1(g("MSX")["T"]) + " s", "tempo mantenuto"))
    if g("LOS"):
        out.append(("Limiti di stabilità", f1(g("LOS")["medie"].get("mxe"), 0) + " %", "escursione massima media"))
    return out


def profilo(prove: dict) -> list[tuple[str, float]]:
    oa = prove.get("OA")
    if not oa:
        return []
    out = []
    for nome, pid in (("Propriocezione", "OC"), ("Vista", "GOA"), ("Vestibolo", "GOC"), ("Movimento visivo", "OPT")):
        p = prove.get(pid)
        if p and p.get("vel"):
            out.append((nome, min(100.0, 100 * oa["vel"] / p["vel"])))
    return out


# ════════════════════════════════════════════════════════════════════
#  GRAFICI (matplotlib)
# ════════════════════════════════════════════════════════════════════

def _plt():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.size": 9, "axes.spines.top": False, "axes.spines.right": False})
    return plt


def _ellisse(ax, e, colore, tratto="-"):
    import numpy as np
    t = np.linspace(0, 2 * np.pi, 120)
    a, b, ang = e.get("a", 0), e.get("b", 0), e.get("ang", 0)
    x = a * np.cos(t) * np.cos(ang) - b * np.sin(t) * np.sin(ang)
    y = a * np.cos(t) * np.sin(ang) + b * np.sin(t) * np.cos(ang)
    ax.plot(x, y, tratto, color=colore, lw=1.8)


def fig_tracciati(prove: dict, ids: list[str], confronto: dict | None = None):
    plt = _plt()
    ids = [i for i in ids if prove.get(i, {}).get("serie")]
    if not ids:
        return None
    n = len(ids)
    col = min(n, 4)
    righe = (n + col - 1) // col
    fig, axs = plt.subplots(righe, col, figsize=(3.0 * col, 3.0 * righe), squeeze=False)
    lim = 10
    for i in ids:
        for P in (prove, confronto or {}):
            for q in P.get(i, {}).get("serie", []):
                lim = max(lim, abs(q["x"]), abs(q["y"]))
    lim = (int(lim * 1.1 / 5) + 1) * 5
    for k, i in enumerate(ids):
        ax = axs[k // col][k % col]
        c = COLORE.get(i, VIOLA)
        if confronto and confronto.get(i, {}).get("serie"):
            s2 = confronto[i]["serie"]
            ax.plot([q["x"] for q in s2], [q["y"] for q in s2], color="#999", lw=0.6, alpha=.6)
            _ellisse(ax, confronto[i].get("ellisse", {}), "#888", "--")
        s = prove[i]["serie"]
        ax.plot([q["x"] for q in s], [q["y"] for q in s], color=c, lw=0.6, alpha=.75)
        _ellisse(ax, prove[i].get("ellisse", {}), c)
        ax.set_xlim(-lim, lim); ax.set_ylim(-lim, lim); ax.set_aspect("equal")
        ax.set_title(f"{PROVE.get(i, i)}\n{f1(prove[i].get('area'), 0)} mm²", fontsize=9)
        ax.axhline(0, color="#ddd", lw=.6); ax.axvline(0, color="#ddd", lw=.6)
    for k in range(n, righe * col):
        axs[k // col][k % col].axis("off")
    fig.tight_layout()
    return fig


def fig_profilo(prof, prof2=None):
    plt = _plt()
    if not prof:
        return None
    fig, ax = plt.subplots(figsize=(5.4, 0.55 * len(prof) + 0.6))
    nomi = [p[0] for p in prof][::-1]
    val = [p[1] for p in prof][::-1]
    colori = ["#B5461F" if v < 60 else VIOLA if v < 80 else "#1D6B44" for v in val]
    ax.barh(nomi, val, color=colori)
    if prof2:
        d = dict(prof2)
        for j, n in enumerate(nomi):
            if n in d:
                ax.plot([d[n], d[n]], [j - .4, j + .4], "k--", lw=1.4)
    for j, v in enumerate(val):
        ax.text(min(v, 100) + 1, j, f1(v, 0) + "%", va="center")
    ax.set_xlim(0, 110); ax.set_xlabel("% (100 = nessun peggioramento rispetto a occhi aperti)")
    fig.tight_layout()
    return fig


def fig_los(los, los2=None):
    plt = _plt()
    import numpy as np
    fig = plt.figure(figsize=(4.2, 4.2))
    ax = fig.add_subplot(111, projection="polar")
    def poli(L, stile, colore):
        dd = sorted(L["dirs"], key=lambda d: d["ang"])
        a = [np.deg2rad(d["ang"]) for d in dd] + [np.deg2rad(dd[0]["ang"])]
        r = [min(d.get("mxe") or 0, 130) for d in dd] + [min(dd[0].get("mxe") or 0, 130)]
        ax.plot(a, r, stile, color=colore, lw=2)
    if los2:
        poli(los2, "--", "#888")
    poli(los, "-", "#2D5BA8")
    ax.set_xticks([np.deg2rad(k) for k in LOS_NOMI]); ax.set_xticklabels([LOS_NOMI[k] for k in LOS_NOMI], fontsize=8)
    ax.set_ylim(0, 120); ax.set_yticks([50, 100]); ax.set_yticklabels(["50%", "100%"], fontsize=7)
    fig.tight_layout()
    return fig


def _png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=160, bbox_inches="tight")
    _plt().close(fig)
    return buf.getvalue()


# ════════════════════════════════════════════════════════════════════
#  TABELLE
# ════════════════════════════════════════════════════════════════════

def tabella_parametri(prove: dict):
    import pandas as pd
    ids = [i for i in PROVE if i in prove and i != "LOS"]
    dati = {}
    for nome, k, d, peggio in PARAMETRI:
        riga = {}
        for i in ids:
            v = prove[i].get(k)
            txt = f1(v, d)
            rif = RIFERIMENTO.get(i)
            if peggio and rif in prove and v is not None and prove[rif].get(k):
                pc = 100 * (v / prove[rif][k] - 1)
                txt += f" ({'+' if pc >= 0 else ''}{f1(pc, 0)}%)"
            riga[PROVE[i]] = txt
        dati[nome] = riga
    return pd.DataFrame(dati).T


def tabella_los(los):
    import pandas as pd
    righe = [{"Direzione": LOS_NOMI.get(d["ang"], d["ang"]), "Reazione (s)": f1(d.get("rt"), 2),
              "Velocità (mm/s)": f1(d.get("vel"), 0), "Punto finale (%)": f1(d.get("epe"), 0),
              "Escursione max (%)": f1(d.get("mxe"), 0), "Controllo direzionale (%)": f1(d.get("dcl"), 0)}
             for d in los["dirs"]]
    m = los["medie"]
    righe.append({"Direzione": "Media", "Reazione (s)": f1(m.get("rt"), 2), "Velocità (mm/s)": f1(m.get("vel"), 0),
                  "Punto finale (%)": f1(m.get("epe"), 0), "Escursione max (%)": f1(m.get("mxe"), 0),
                  "Controllo direzionale (%)": f1(m.get("dcl"), 0)})
    return pd.DataFrame(righe).set_index("Direzione")


def tabella_confronto(A: dict, B: dict):
    import pandas as pd
    ids = [i for i in PROVE if i in A and i in B and i != "LOS"]
    dati = {}
    for nome, k, d, peggio in PARAMETRI:
        if not peggio:
            continue
        riga = {}
        for i in ids:
            a, b = A[i].get(k), B[i].get(k)
            if a and b is not None:
                pc = 100 * (b / a - 1)
                riga[PROVE[i]] = f"{f1(a, d)} → {f1(b, d)} ({'+' if pc >= 0 else ''}{f1(pc, 0)}%)"
        dati[nome] = riga
    return pd.DataFrame(dati).T


# ════════════════════════════════════════════════════════════════════
#  REFERTO WORD
# ════════════════════════════════════════════════════════════════════

def referto_docx(paziente: str, esame: dict) -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    P = esame["payload"]
    prove = P.get("prove", {})
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.8)
    st_ = doc.styles["Normal"]; st_.font.name = "Calibri"; st_.font.size = Pt(10)
    h = sec.header.paragraphs[0]
    r = h.add_run("STUDIO THE ORGANISM  ·  Metodo PNEV"); r.bold = True; r.font.color.rgb = RGBColor(0x1D, 0x6B, 0x44)
    f = sec.footer.paragraphs[0]
    f.add_run("Pagani · Via De Rosa 46   |   Sant'Agnello · Via Balsamo 19   |   Avellino · Via Renzi 40   ·   "
              "WhatsApp 391 3598767 · apstheorganism@gmail.com · www.pnev.it").font.size = Pt(7)

    t = doc.add_paragraph(); r = t.add_run("Esame stabilometrico"); r.bold = True; r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x1D, 0x6B, 0x44)
    doc.add_paragraph(f"Paziente: {paziente}    Data: {_data(esame['data_esame'])}    "
                      f"Durata delle prove: {str(P.get('durata', '')).replace('.', ',')} s")
    ind = indici(prove)
    if ind:
        doc.add_heading("Indici di sintesi", level=2)
        tb = doc.add_table(rows=0, cols=3); tb.style = "Light Grid Accent 1"
        for n, v, s in ind:
            c = tb.add_row().cells; c[0].text = n; c[1].text = v; c[2].text = s
    commento = esame.get("commento") or P.get("commento") or ""
    if commento:
        doc.add_heading("Commento", level=2)
        doc.add_paragraph(commento)
    fig = fig_tracciati(prove, [i for i in PROVE if i in prove])
    if fig:
        doc.add_heading("Tracciati (statokinesigrammi con ellisse al 95%)", level=2)
        doc.add_picture(io.BytesIO(_png(fig)), width=Cm(17))
    prof = profilo(prove)
    if prof:
        doc.add_heading("Profilo sensoriale", level=2)
        doc.add_picture(io.BytesIO(_png(fig_profilo(prof))), width=Cm(12))
    if prove.get("LOS"):
        doc.add_heading("Limiti di stabilità", level=2)
        doc.add_picture(io.BytesIO(_png(fig_los(prove["LOS"]))), width=Cm(9))
        df = tabella_los(prove["LOS"])
        tb = doc.add_table(rows=1, cols=len(df.columns) + 1); tb.style = "Light Grid Accent 1"
        tb.rows[0].cells[0].text = "Direzione"
        for j, c in enumerate(df.columns):
            tb.rows[0].cells[j + 1].text = c
        for idx, row in df.iterrows():
            cc = tb.add_row().cells; cc[0].text = str(idx)
            for j, v in enumerate(row):
                cc[j + 1].text = str(v)
    df = tabella_parametri(prove)
    if not df.empty:
        doc.add_heading("Parametri", level=2)
        tb = doc.add_table(rows=1, cols=len(df.columns) + 1); tb.style = "Light Grid Accent 1"
        tb.rows[0].cells[0].text = "Parametro"
        for j, c in enumerate(df.columns):
            tb.rows[0].cells[j + 1].text = c
        for idx, row in df.iterrows():
            cc = tb.add_row().cells; cc[0].text = str(idx)
            for j, v in enumerate(row):
                cc[j + 1].text = "" if v is None or v != v else str(v)
        for row in tb.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(7.5)
    n = doc.add_paragraph()
    r = n.add_run("Esame eseguito con pedana Wii Balance Board: strumento non certificato come dispositivo "
                  "medico. I valori servono al confronto nello stesso paziente nel tempo, con la stessa "
                  "procedura; vanno integrati con l'esame clinico.")
    r.italic = True; r.font.size = Pt(8)
    doc.add_paragraph("\n\nIl professionista ______________________________")
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


# ════════════════════════════════════════════════════════════════════
#  INTERFACCIA
# ════════════════════════════════════════════════════════════════════

def _leggi_file(file) -> list[dict]:
    try:
        dati = json.loads(file.getvalue().decode("utf-8"))
    except Exception:
        return []
    voci = dati if isinstance(dati, list) else [dati]
    return [v for v in voci if isinstance(v, dict) and v.get("tipo") in ("esame", "bio")]


def _scheda_carica(conn, paz_id, nome_paz):
    st.markdown("Carica il file creato dalla pagina della pedana con **«Esporta per il gestionale»**. "
                "Va bene anche la **copia di sicurezza** dell'archivio: prendo solo gli esami di questo paziente.")
    files = st.file_uploader("File degli esami (.json)", type=["json"], accept_multiple_files=True,
                             key=f"stab_upl_{paz_id}")
    if not files:
        return
    voci = []
    for f in files:
        voci += _leggi_file(f)
    if not voci:
        st.error("Nei file non ho trovato esami della pagina Stabilometria.")
        return
    gia = {int(r["origine_id"]) for r in _righe(conn, "SELECT origine_id FROM stabilometria_esami "
                                                     "WHERE paziente_id=%s AND origine_id IS NOT NULL", (int(paz_id),))}
    mio = _norm(nome_paz)
    st.markdown(f"**{len(voci)}** voci trovate. Spunta quelle da salvare nella scheda di **{nome_paz}**.")
    scelte = []
    for i, v in enumerate(voci):
        stesso = _norm(v.get("paziente", "")) == mio
        doppio = v.get("id") in gia
        etichetta = (f"{_data(v.get('data'))} · {'Esame' if v.get('tipo') == 'esame' else 'Biofeedback'} · "
                     f"«{v.get('paziente') or 'senza nome'}»"
                     + (f" · prove: {', '.join(PROVE.get(k, k) for k in (v.get('prove') or {}))}" if v.get("tipo") == "esame"
                        else f" · {v.get('gioco', '')}"))
        if doppio:
            st.caption("✓ già salvato: " + etichetta)
            continue
        if st.checkbox(etichetta + ("" if stesso else "  ⚠️ nome diverso dal paziente in lavorazione"),
                       value=stesso, key=f"stab_sel_{paz_id}_{i}_{v.get('id')}"):
            scelte.append(v)
    if scelte and st.button(f"💾 Salva {len(scelte)} nella scheda", type="primary", key=f"stab_salva_{paz_id}"):
        errori = 0
        for v in scelte:
            err = _esegui(conn, "INSERT INTO stabilometria_esami (paziente_id, data_esame, tipo, origine_id, "
                                "nome_su_file, payload, commento) VALUES (%s,%s,%s,%s,%s,%s,%s) "
                                "ON CONFLICT (paziente_id, origine_id) DO NOTHING",
                          (int(paz_id), v.get("data") or None, v.get("tipo"), v.get("id"), v.get("paziente"),
                           json.dumps(v), v.get("commento") or v.get("riassunto") or ""))
            errori += bool(err)
        if errori:
            st.error(f"{errori} voci non salvate. Riprova.")
        else:
            st.success(f"{len(scelte)} voci salvate nella scheda di {nome_paz}.")
            st.rerun()


def _scheda_dettaglio(conn, paz_id, nome_paz, esami):
    esami_t = [e for e in esami if e["tipo"] == "esame"]
    bio = [e for e in esami if e["tipo"] == "bio"]
    if not esami_t and not bio:
        st.info("Nessun esame salvato per questo paziente. Caricali dalla scheda «Carica esami».")
        return
    if esami_t:
        opz = {f"{_data(e['data_esame'])} · {', '.join(PROVE.get(k, k) for k in e['payload'].get('prove', {}))}": e
               for e in reversed(esami_t)}
        e = opz[st.selectbox("Esame", list(opz), key=f"stab_det_{paz_id}")]
        prove = e["payload"].get("prove", {})
        ind = indici(prove)
        if ind:
            cols = st.columns(min(len(ind), 4))
            for k, (n, v, s) in enumerate(ind):
                cols[k % len(cols)].metric(n, v, help=s)
        fig = fig_tracciati(prove, list(PROVE))
        if fig:
            st.pyplot(fig)
        c1, c2 = st.columns(2)
        prof = profilo(prove)
        if prof:
            with c1:
                st.markdown("**Profilo sensoriale**")
                st.pyplot(fig_profilo(prof))
        if prove.get("LOS"):
            with c2:
                st.markdown("**Limiti di stabilità**")
                st.pyplot(fig_los(prove["LOS"]))
            st.dataframe(tabella_los(prove["LOS"]), use_container_width=True)
        st.markdown("**Parametri**")
        st.dataframe(tabella_parametri(prove), use_container_width=True)
        st.markdown("**Commento**")
        nuovo = st.text_area("Commento", value=e.get("commento") or "", height=170,
                             key=f"stab_comm_{e['id']}", label_visibility="collapsed")
        b1, b2, b3 = st.columns(3)
        if b1.button("💾 Salva il commento", key=f"stab_scomm_{e['id']}"):
            err = _esegui(conn, "UPDATE stabilometria_esami SET commento=%s WHERE id=%s", (nuovo, e["id"]))
            st.error(err) if err else st.success("Commento salvato.")
        try:
            e2 = dict(e, commento=nuovo)
            b2.download_button("📄 Referto Word", data=referto_docx(nome_paz, e2),
                               file_name=f"Stabilometria_{nome_paz.replace(' ', '_')}_{e['data_esame']}.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               key=f"stab_doc_{e['id']}")
        except Exception as ex:
            b2.caption(f"Referto non disponibile: {ex}")
        if b3.button("🗑️ Elimina questo esame", key=f"stab_del_{e['id']}"):
            st.session_state[f"stab_conf_del_{e['id']}"] = True
        if st.session_state.get(f"stab_conf_del_{e['id']}"):
            st.warning("Eliminare definitivamente questo esame?")
            if st.button("Sì, elimina", key=f"stab_del2_{e['id']}"):
                _esegui(conn, "DELETE FROM stabilometria_esami WHERE id=%s", (e["id"],))
                st.session_state.pop(f"stab_conf_del_{e['id']}", None)
                st.rerun()
    if bio:
        st.markdown("**Sedute di biofeedback**")
        import pandas as pd
        st.dataframe(pd.DataFrame([{"Data": _data(b["data_esame"]), "Esercizio": b["payload"].get("gioco", ""),
                                    "Risultato": b["payload"].get("riassunto", "")} for b in bio]),
                     use_container_width=True, hide_index=True)


def _scheda_confronto(conn, paz_id, esami):
    import pandas as pd
    esami_t = [e for e in esami if e["tipo"] == "esame"]
    if len(esami_t) < 2:
        st.info("Servono almeno due esami salvati per il confronto e l'andamento.")
        return
    st.markdown("**Andamento nel tempo**")
    righe = []
    for e in esami_t:
        p = e["payload"].get("prove", {})
        righe.append({"Data": pd.to_datetime(e["data_esame"]),
                      "Area OA (mm²)": p.get("OA", {}).get("area"), "Area OC (mm²)": p.get("OC", {}).get("area"),
                      "Velocità OA (mm/s)": p.get("OA", {}).get("vel"), "Velocità OC (mm/s)": p.get("OC", {}).get("vel")})
    df = pd.DataFrame(righe).set_index("Data")
    c1, c2 = st.columns(2)
    c1.line_chart(df[["Area OA (mm²)", "Area OC (mm²)"]])
    c2.line_chart(df[["Velocità OA (mm/s)", "Velocità OC (mm/s)"]])
    st.markdown("**Confronto tra due esami**")
    opz = {f"{_data(e['data_esame'])} (n. {e['id']})": e for e in esami_t}
    nomi = list(opz)
    c1, c2 = st.columns(2)
    a = c1.selectbox("Prima", nomi, index=0, key=f"stab_cA_{paz_id}")
    b = c2.selectbox("Dopo", nomi, index=len(nomi) - 1, key=f"stab_cB_{paz_id}")
    if a == b:
        st.caption("Scegli due esami diversi.")
        return
    A, B = opz[a]["payload"].get("prove", {}), opz[b]["payload"].get("prove", {})
    comuni = [i for i in PROVE if i in A and i in B and i != "LOS"]
    fig = fig_tracciati(B, comuni, confronto=A)
    if fig:
        st.caption("Grigio tratteggiato = esame «prima»; colorato = esame «dopo».")
        st.pyplot(fig)
    iA, iB = dict((n, v) for n, v, _ in indici(A)), dict((n, v) for n, v, _ in indici(B))
    comuni_i = [n for n in iA if n in iB]
    if comuni_i:
        st.dataframe(pd.DataFrame({"Prima": [iA[n] for n in comuni_i], "Dopo": [iB[n] for n in comuni_i]},
                                  index=comuni_i), use_container_width=True)
    c1, c2 = st.columns(2)
    pA, pB = profilo(A), profilo(B)
    if pB:
        with c1:
            st.markdown("**Profilo sensoriale** (barre = dopo, tratteggio = prima)")
            st.pyplot(fig_profilo(pB, pA))
    if A.get("LOS") and B.get("LOS"):
        with c2:
            st.markdown("**Limiti di stabilità** (tratteggio = prima)")
            st.pyplot(fig_los(B["LOS"], A["LOS"]))
    st.markdown("**Parametri: prima → dopo**")
    st.dataframe(tabella_confronto(A, B), use_container_width=True)


def _scheda_pagina():
    import os
    st.markdown("La pagina **Stabilometria PNEV** si apre in **Google Chrome** sul computer collegata alla "
                "pedana (Bluetooth). Lì si fanno gli esami e il biofeedback; poi con «Esporta per il "
                "gestionale» il file si carica qui.")
    percorsi = [PAGINA_PEDANA, os.path.join(os.path.dirname(__file__), "..", PAGINA_PEDANA)]
    for p in percorsi:
        if os.path.exists(p):
            with open(p, "rb") as f:
                st.download_button("⬇️ Scarica la pagina della pedana", f.read(),
                                   file_name="Stabilometria_PNEV.html", mime="text/html", type="primary")
            break
    else:
        st.info(f"Per scaricarla da qui, carica il file della pagina nel repository come `{PAGINA_PEDANA}`.")
    st.markdown("""
**Prima volta sul computer**
1. Abbina la pedana a Windows con Dolphin (una volta sola): Controller → Balance Board reale → Scansione continua → tastino rosso.
2. Apri la pagina in Chrome, accendi la pedana dal pulsante davanti, premi «Collega la pedana».
""")


def _sessioni_vecchie(conn, paz_id):
    righe = _righe(conn, "SELECT data_sessione, condizione, peso_sx_pct, peso_dx_pct, oscillazione_mm, note "
                         "FROM postura_wiifit_sessioni WHERE paziente_id=%s ORDER BY creato_il DESC LIMIT 30",
                   (int(paz_id),))
    if righe:
        import pandas as pd
        with st.expander(f"Sessioni del vecchio modulo Postura ({len(righe)})"):
            st.dataframe(pd.DataFrame(righe), use_container_width=True, hide_index=True)


def render_stabilometria(conn=None, paz_id=None, paziente=None) -> None:
    st.header("🩶 Stabilometria — pedana Wii")
    st.caption("Esami fatti con la pagina Stabilometria PNEV e salvati qui per paziente. Strumento non "
               "certificato come dispositivo medico: valori per il confronto nel tempo.")
    if conn is None:
        st.info("Connessione al database non disponibile.")
        return
    if not paz_id:
        st.info("Seleziona un paziente qui sopra.")
        return
    _assicura_tabella(conn)
    nome_paz = _nome_paziente(conn, paz_id) or f"Paziente {paz_id}"
    esami = _esami(conn, paz_id)
    n_es = sum(1 for e in esami if e["tipo"] == "esame")
    t1, t2, t3, t4 = st.tabs([f"📋 Esami ({n_es})", "📈 Confronto e andamento", "📥 Carica esami",
                              "⬇️ Pagina della pedana"])
    with t1:
        _scheda_dettaglio(conn, paz_id, nome_paz, esami)
        _sessioni_vecchie(conn, paz_id)
    with t2:
        _scheda_confronto(conn, paz_id, esami)
    with t3:
        _scheda_carica(conn, paz_id, nome_paz)
    with t4:
        _scheda_pagina()
