# -*- coding: utf-8 -*-
"""
╔══════════════════════════════════════════════════════════════════════╗
║  RELAZIONE IN WORD — la stessa relazione, ma modificabile            ║
║                                                                      ║
║  Il gestionale produce relazioni in PDF: impaginate, con carta       ║
║  intestata, timbro e firma. Vanno benissimo per consegnare, ma un    ║
║  PDF non si corregge: se dopo la stampa ci si accorge di una frase   ║
║  da cambiare bisogna tornare nel modulo, rigenerare tutto e          ║
║  ristampare — e le integrazioni scritte a mano dal professionista    ║
║  fuori dai campi previsti non si possono fare affatto.               ║
║                                                                      ║
║  Questo modulo genera lo STESSO contenuto in .docx: stesso testo,    ║
║  stesse sezioni, stessa intestazione, ma apribile in Word, Pages,    ║
║  LibreOffice o Google Documenti e modificabile riga per riga.        ║
║                                                                      ║
║  Il PDF resta: si usa quello per consegnare, il Word quando la       ║
║  relazione va rivista o integrata prima della firma.                 ║
║                                                                      ║
║  Nota sul timbro: nel Word non viene inserito. Un documento ancora   ║
║  modificabile non e' un documento firmato, e un timbro su una bozza  ║
║  che chiunque puo' cambiare e' peggio di nessun timbro. Il timbro    ║
║  compare nel PDF, che e' la versione da consegnare.                  ║
╚══════════════════════════════════════════════════════════════════════╝
"""
from __future__ import annotations

import io

MIME_DOCX = ("application/vnd.openxmlformats-officedocument."
             "wordprocessingml.document")

VERDE = "1D6B44"
GRIGIO = "888780"

INDIRIZZO = ("Via De Rosa 46, Pagani (SA) · Via Tino di Camaino 23, Napoli · "
             "Studio Della Ragione, Via Balsamo 19, Sant'Agnello (NA)")
CONTATTI = "Tel. 0815152334  |  Cell. 3921873914  |  apstheorganism@gmail.com"


def _rgb(hexstr):
    from docx.shared import RGBColor
    return RGBColor.from_string(hexstr)


def _dati_studio():
    """Indirizzo e contatti salvati in «Intestazione dello studio», così il
    Word riporta la stessa intestazione del PDF invece delle costanti."""
    try:
        from modules.pdf_templates import _intestazione_studio
        return _intestazione_studio() or {}
    except Exception:
        return {}


def genera_docx_carta_intestata(professionista: str, titolo: str,
                                paziente: str, data: str, titolo_doc: str,
                                corpo_testo: str = "",
                                includi_intro_pnev: bool = True) -> bytes:
    """Stessa firma di pdf_templates.genera_carta_intestata, senza timbro.

    Le righe che iniziano con ### diventano titoli di sezione, come nel PDF.
    """
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    for sez in doc.sections:
        sez.top_margin = Cm(2.0)
        sez.bottom_margin = Cm(2.0)
        sez.left_margin = Cm(2.2)
        sez.right_margin = Cm(2.2)

    normale = doc.styles["Normal"]
    normale.font.name = "Calibri"
    normale.font.size = Pt(11)

    # ── Intestazione ──────────────────────────────────────────────────
    p = doc.add_paragraph()
    r = p.add_run(professionista or "")
    r.bold = True
    r.font.size = Pt(12)
    if titolo:
        p2 = doc.add_paragraph()
        r2 = p2.add_run(titolo)
        r2.font.size = Pt(9)
        r2.font.color.rgb = _rgb(GRIGIO)

    _d = _dati_studio()
    _indirizzo = " · ".join(
        r.strip() for r in str(_d.get("indirizzo") or "").splitlines() if r.strip()
    ) or INDIRIZZO
    _contatti = (str(_d.get("contatti") or "").strip()
                 or str(_d.get("telefono") or "").strip()
                 or CONTATTI)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(_indirizzo + "\n" + _contatti)
    r3.font.size = Pt(7.5)
    r3.font.color.rgb = _rgb(GRIGIO)

    # ── Titolo del documento ──────────────────────────────────────────
    pt = doc.add_paragraph()
    pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = pt.add_run((titolo_doc or "").upper())
    rt.bold = True
    rt.font.size = Pt(14)
    rt.font.color.rgb = _rgb(VERDE)

    # ── Paziente e data ───────────────────────────────────────────────
    tab = doc.add_table(rows=1, cols=2)
    tab.autofit = True
    c1 = tab.cell(0, 0).paragraphs[0]
    c1.add_run("Paziente: ").bold = True
    c1.add_run(paziente or "")
    c2 = tab.cell(0, 1).paragraphs[0]
    c2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    c2.add_run("Data: ").bold = True
    c2.add_run(data or "")
    doc.add_paragraph()

    # ── Corpo ─────────────────────────────────────────────────────────
    testo = corpo_testo or ""
    if includi_intro_pnev:
        try:
            from modules.relazione_testi import intro_pnev, bibliografia
            testo = intro_pnev() + "\n" + testo + "\n\n" + bibliografia()
        except Exception:
            pass

    for riga in testo.split("\n"):
        if riga.startswith("###"):
            ph = doc.add_paragraph()
            rh = ph.add_run(riga.replace("###", "").strip())
            rh.bold = True
            rh.font.size = Pt(12)
            rh.font.color.rgb = _rgb(VERDE)
        elif riga.strip():
            pp = doc.add_paragraph(riga)
            pp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            doc.add_paragraph()

    # ── Firma ─────────────────────────────────────────────────────────
    doc.add_paragraph()
    tf = doc.add_table(rows=2, cols=2)
    tf.cell(0, 0).paragraphs[0].add_run("_" * 28)
    pfd = tf.cell(0, 1).paragraphs[0]
    pfd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pfd.add_run("_" * 28)
    rf1 = tf.cell(1, 0).paragraphs[0].add_run("Firma e timbro")
    rf1.font.size = Pt(8)
    rf1.font.color.rgb = _rgb(GRIGIO)
    pfd2 = tf.cell(1, 1).paragraphs[0]
    pfd2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rf2 = pfd2.add_run(professionista or "")
    rf2.font.size = Pt(8)
    rf2.font.color.rgb = _rgb(GRIGIO)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def bottone_word(st, etichetta: str, nome_file: str, key: str, **kwargs) -> None:
    """Bottone di download Word accanto a quello PDF.

    Gli argomenti di genera_docx_carta_intestata si passano come kwargs.
    Se python-docx non e' disponibile lo dice invece di sparire in
    silenzio: un bottone che non compare sembra una funzione che non
    esiste."""
    try:
        dati = genera_docx_carta_intestata(**kwargs)
    except ImportError:
        st.caption("Export Word non disponibile: manca python-docx fra le dipendenze.")
        return
    except Exception as e:
        st.caption(f"Export Word non riuscito: {e}")
        return
    st.download_button(etichetta, data=dati, file_name=nome_file,
                       mime=MIME_DOCX, key=key)
